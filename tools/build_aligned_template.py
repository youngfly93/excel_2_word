"""
Build an aligned docx template with docxtpl placeholders so that static
structure matches the reference final report as much as practical, while
keeping variable regions driven by Excel data.

Output: templates/aligned_template.docx

Run:
  python3 tools/build_aligned_template.py
"""

from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from pathlib import Path


def set_a4_margins(doc: Document, left=3.175, right=3.175, top=2.54, bottom=2.54):
    sec = doc.sections[0]
    # A4 portrait
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.left_margin = Cm(left)
    sec.right_margin = Cm(right)
    sec.top_margin = Cm(top)
    sec.bottom_margin = Cm(bottom)


def add_title(doc: Document, text: str, size=20, bold=True):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_heading(doc: Document, text: str, level: int = 1):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(14 if level == 1 else 12)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT


def add_paragraph(doc: Document, text: str, size=11):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(size)
    return p


def add_key_value_table(doc: Document, rows: list[tuple[str, str]]):
    table = doc.add_table(rows=len(rows), cols=2)
    table.style = "Table Grid"
    for i, (k, v) in enumerate(rows):
        table.cell(i, 0).text = k
        table.cell(i, 1).text = v
    return table


def add_table_with_headers(doc: Document, headers: list[str]):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for j, h in enumerate(headers):
        table.cell(0, j).text = h
    return table


def build_variants_table(doc: Document):
    headers = [
        "基因",
        "cHGVS",
        "蛋白改变",
        "变异类型",
        "频率(%)",
        "外显子",
        "相关药物",
        "证据等级",
    ]
    table = add_table_with_headers(doc, headers)
    # loop start row
    loop_start = table.add_row().cells
    loop_start[0].text = "{% for row in variants %}"
    # data row
    data = table.add_row().cells
    data[0].text = "{{ row.Gene_Symbol }}"
    data[1].text = "{{ row.cHGVS }}"
    data[2].text = "{{ row.pHGVS_S or row.pHGVS_A }}"
    data[3].text = "{{ row.Function }}"
    data[4].text = "{{ row['Freq(%)'] }}"
    data[5].text = "{{ row.ExIn_ID }}"
    data[6].text = "{{ row.Drug }}"
    data[7].text = "{{ row.Evidence_level }}"
    # loop end row
    loop_end = table.add_row().cells
    loop_end[0].text = "{% endfor %}"


def build_chemotherapy_table(doc: Document):
    headers = ["药物名称", "相关基因", "药物适应情况"]
    table = add_table_with_headers(doc, headers)
    # loop start
    s = table.add_row().cells
    s[0].text = "{% for row in chemotherapy %}"
    # data row (use safe keys, blanks if missing)
    d = table.add_row().cells
    d[0].text = "{{ row.药物 or row.Drug or '' }}"
    d[1].text = "{{ row.检测基因 or row.Gene or '' }}"
    d[2].text = "{{ row.药物适应情况 or row.检测结果 or '' }}"
    # loop end
    e = table.add_row().cells
    e[0].text = "{% endfor %}"


def build_tmb_msi_table(doc: Document):
    headers = ["TMB/MSI/其它生物标志物检测结果", "TMB/MSI/其它生物标志物检测结果", "用药提示"]
    table = add_table_with_headers(doc, headers)
    row1 = table.add_row().cells
    row1[0].text = "TMB"
    # 使用统一变量名，兼容mapping.yaml
    row1[1].text = "{{ tmb_value }} {{ tmb_unit }}"
    row1[2].text = "{{ immuno_tips }}"
    row2 = table.add_row().cells
    row2[0].text = "MSI状态"
    row2[1].text = "{{ msi_status }}"
    row2[2].text = ""


def main():
    out = Path("templates/aligned_template.docx")
    out.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    set_a4_margins(doc)

    # 封面
    add_title(doc, "检测报告", size=26)
    add_paragraph(doc, "")
    add_heading(doc, "患者信息", level=1)
    add_key_value_table(
        doc,
        [
            ("姓    名：", "{{ patient_name|default('') }}"),
            ("样本编号：", "{{ sample_id|default('') }}"),
            ("性    别：", "{{ gender|default('') }}"),
            ("年    龄：", "{{ age|default('') }}"),
            ("病 理 号：", "{{ pathology_id|default('-') }}"),
            ("送检医院：", "{{ hospital|default('-') }}"),
            ("科    室：", "{{ department|default('-') }}"),
            ("报告日期：", "{{ report_date|default('') }}"),
        ],
    )

    add_paragraph(doc, "")
    add_heading(doc, "报告导读", level=1)
    add_paragraph(doc, "报告第一部分：检测报告的基本信息。包括患者及样本信息、检测内容。")
    add_paragraph(doc, "报告第二部分：检测结果（靶向、免疫、化疗等综合结果），为报告关键信息。")
    add_paragraph(doc, "报告第三部分：基因变异及药物解析，并包含本报告的阅读说明。")
    add_paragraph(doc, "报告第四部分：附录。提供诊疗知识、信号通路、基因列表、参考文献等。")
    add_paragraph(doc, "咨询电话：022-87190699。")

    add_paragraph(doc, "")
    add_heading(doc, "目    录", level=1)

    add_paragraph(doc, "")
    add_heading(doc, "第一部分：基本信息", level=1)
    add_heading(doc, "患者信息", level=2)
    # 再次展示关键信息（可视化更清楚）
    add_key_value_table(
        doc,
        [
            ("姓    名：", "{{ patient_name|default('') }}"),
            ("样本编号：", "{{ sample_id|default('') }}"),
            ("性    别：", "{{ gender|default('') }}"),
            ("年    龄：", "{{ age|default('') }}"),
            ("送检医院：", "{{ hospital|default('-') }}"),
        ],
    )
    add_paragraph(doc, "附注：以上受检者/样本信息来自送检时提供信息，本报告不对其准确性负责。")

    add_heading(doc, "检测内容", level=2)
    tbl = add_table_with_headers(doc, ["检测基因", "检测内容", "检测结果"])
    # 留空行作为示例
    tbl.add_row()

    add_paragraph(doc, "")
    add_heading(doc, "第二部分：检测结果", level=1)
    add_heading(doc, "TMB/MSI/其它生物标志物结果", level=2)
    build_tmb_msi_table(doc)

    add_paragraph(doc, "")
    add_heading(doc, "化疗药物检测结果", level=2)
    build_chemotherapy_table(doc)

    add_paragraph(doc, "")
    add_heading(doc, "基因变异明细", level=2)
    build_variants_table(doc)

    add_paragraph(doc, "")
    add_heading(doc, "第三部分：基因变异及药物解析", level=1)
    add_paragraph(doc, "（此处为解析与阅读说明静态内容，可后续补充）")

    add_paragraph(doc, "")
    add_heading(doc, "第四部分：附录", level=1)
    add_paragraph(doc, "（此处为诊疗知识/信号通路/基因列表/参考文献等静态内容骨架）")

    out.write_bytes(b"")  # ensure path exists
    doc.save(str(out))
    print(f"Written template: {out}")


if __name__ == "__main__":
    main()
