#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
从手工终版报告中提取化疗药物详细解析表的结构
识别表格11-44的药物列表和表格模式
"""

from docx import Document
import json

def extract_drug_tables(docx_path):
    """提取化疗药物详细解析表"""

    doc = Document(docx_path)

    print(f"📊 分析手工终版报告: {docx_path}")
    print(f"总表格数: {len(doc.tables)}\n")

    drug_tables = []

    # 分析表格11-44（索引10-43）
    for i in range(10, min(44, len(doc.tables))):
        table = doc.tables[i]

        # 提取表格信息
        table_info = {
            "index": i + 1,
            "rows": len(table.rows),
            "columns": len(table.columns),
            "drug_name": None,
            "structure": [],
            "sample_data": []
        }

        # 提取药物名称（通常在第一行第一列或表格标题附近）
        if len(table.rows) > 0:
            first_cell = table.rows[0].cells[0].text.strip()
            # 药物名称通常是表格的第一个单元格
            if first_cell and len(first_cell) < 50:
                table_info["drug_name"] = first_cell

        # 提取表头结构
        if len(table.rows) > 0:
            header_row = table.rows[0]
            headers = [cell.text.strip() for cell in header_row.cells]
            table_info["structure"] = headers

        # 提取前3行数据作为样本
        for j, row in enumerate(table.rows[:3]):
            row_data = [cell.text.strip() for cell in row.cells]
            table_info["sample_data"].append(row_data)

        # 判断是否为化疗药物表格（6列结构）
        if table_info["columns"] == 6:
            drug_tables.append(table_info)
            print(f"表格{i+1}: {table_info['drug_name']}")
            print(f"  结构: {table_info['rows']}行 x {table_info['columns']}列")
            print(f"  表头: {', '.join(table_info['structure'])}")
            print()

    return drug_tables

def extract_drug_list(drug_tables):
    """提取药物列表"""

    drugs = []
    for table_info in drug_tables:
        drug_name = table_info.get("drug_name")
        if drug_name:
            drugs.append({
                "name": drug_name,
                "table_index": table_info["index"],
                "rows": table_info["rows"],
                "columns": table_info["columns"]
            })

    return drugs

def main():
    import sys

    if len(sys.argv) < 2:
        print("用法: python extract_drug_tables.py <手工终版报告.docx>")
        sys.exit(1)

    docx_path = sys.argv[1]

    print("=" * 80)
    print("🔍 提取化疗药物详细解析表")
    print("=" * 80)
    print()

    # 提取药物表格
    drug_tables = extract_drug_tables(docx_path)

    print("=" * 80)
    print(f"✅ 找到 {len(drug_tables)} 个化疗药物详细解析表")
    print("=" * 80)

    # 提取药物列表
    drugs = extract_drug_list(drug_tables)

    print("\n📋 药物列表:")
    for i, drug in enumerate(drugs, 1):
        print(f"  {i}. {drug['name']} (表格{drug['table_index']}, {drug['rows']}行)")

    # 保存结果
    output = {
        "total_drug_tables": len(drug_tables),
        "drug_tables": drug_tables,
        "drug_list": drugs
    }

    output_file = "drug_tables_analysis.json"
    with open(output_file, 'w', encoding='utf-8') as f:
        json.dump(output, f, ensure_ascii=False, indent=2)

    print(f"\n✅ 详细结果已保存到: {output_file}")
    print("=" * 80)

if __name__ == "__main__":
    main()
