#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
验证生成的报告

检查HLA和Variations数据是否正确包含在报告中
"""

from docx import Document

def verify_report(docx_file="data/output/张三---结直肠癌358基因检测-MLB2509307001-终版.docx"):
    """验证生成的报告"""

    print("=" * 80)
    print("验证生成的报告")
    print("=" * 80)

    print(f"\n📂 报告文件: {docx_file}")

    # 打开文档
    doc = Document(docx_file)

    print(f"\n【文档基本信息】")
    print(f"   段落数: {len(doc.paragraphs)}")
    print(f"   表格数: {len(doc.tables)}")

    # 查找HLA相关内容
    print(f"\n【HLA数据验证】")

    hla_found = False
    hla_tables_count = 0

    for i, table in enumerate(doc.tables):
        # 检查表头是否包含HLA相关列名
        if table.rows:
            first_row = table.rows[0]
            header_text = ' '.join(cell.text for cell in first_row.cells).lower()

            if 'hla' in header_text or 'locus' in header_text or '位点' in header_text:
                hla_found = True
                hla_tables_count += 1
                print(f"   ✅ 找到HLA表格 (表格 {i+1})")
                print(f"      表头: {[cell.text for cell in first_row.cells][:5]}")

                # 显示前3行数据
                if len(table.rows) > 1:
                    print(f"      数据行数: {len(table.rows) - 1}")
                    for row_idx in range(1, min(4, len(table.rows))):
                        row_data = [cell.text for cell in table.rows[row_idx].cells]
                        print(f"         Row {row_idx}: {row_data[:5]}")

    if not hla_found:
        print(f"   ⚠️  未找到HLA表格")

    # 查找Variations相关内容
    print(f"\n【Variations数据验证】")

    variations_found = False
    variations_rows = 0

    for i, table in enumerate(doc.tables):
        if table.rows:
            first_row = table.rows[0]
            header_text = ' '.join(cell.text for cell in first_row.cells).lower()

            # 检查是否是变异表格（包含基因、变异、频率等列）
            # 排除Fusion表格（有"基因1"、"基因2"列）
            is_fusion = '基因1' in header_text or 'gene1' in header_text

            has_gene = any(keyword in header_text for keyword in ['gene', '基因', 'symbol'])
            has_variant = any(keyword in header_text for keyword in ['variant', '变异', 'hgvs', 'mutation', '突变位点'])
            has_freq = any(keyword in header_text for keyword in ['freq', '频率', 'af'])

            if has_gene and (has_variant or has_freq) and not is_fusion:
                variations_found = True
                variations_rows = len(table.rows) - 1  # 减去表头
                print(f"   ✅ 找到Variations表格 (表格 {i+1})")
                print(f"      表头: {[cell.text for cell in first_row.cells][:8]}")
                print(f"      数据行数: {variations_rows}")

                # 显示前5行数据
                if len(table.rows) > 1:
                    print(f"\n      前5行数据:")
                    for row_idx in range(1, min(6, len(table.rows))):
                        row_data = [cell.text for cell in table.rows[row_idx].cells]
                        print(f"         Row {row_idx}: {row_data[:8]}")

                break  # 只检查第一个匹配的表格

    if not variations_found:
        print(f"   ⚠️  未找到Variations表格")

    # 统计总结
    print(f"\n【统计总结】")
    print(f"   HLA表格: {'✅ 已包含' if hla_found else '❌ 未找到'}")
    if hla_found:
        print(f"      HLA表格数量: {hla_tables_count}")

    print(f"   Variations表格: {'✅ 已包含' if variations_found else '❌ 未找到'}")
    if variations_found:
        print(f"      变异行数: {variations_rows}")
        print(f"      预期范围: 40-60 行 (过滤后)")

        if 40 <= variations_rows <= 60:
            print(f"      ✅ 行数在预期范围内")
        else:
            print(f"      ⚠️  行数不在预期范围内")

    print("\n" + "=" * 80)
    print("✅ 报告验证完成")
    print("=" * 80)


if __name__ == "__main__":
    verify_report()
