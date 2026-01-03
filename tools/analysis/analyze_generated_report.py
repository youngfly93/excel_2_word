#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
深度分析生成的报告结构
提取所有表格、段落、变量区的详细信息
"""

from docx import Document
import json
import sys

def analyze_docx_structure(docx_path):
    """分析docx文档的完整结构"""

    doc = Document(docx_path)

    analysis = {
        "file": docx_path,
        "metadata": {
            "paragraphs_count": len(doc.paragraphs),
            "tables_count": len(doc.tables),
            "sections_count": len(doc.sections)
        },
        "paragraphs": [],
        "tables": [],
        "potential_variables": []
    }

    # 分析段落
    print(f"\n📄 分析段落结构...")
    for i, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if text:  # 只记录非空段落
            para_info = {
                "index": i,
                "text": text[:200],  # 限制长度
                "full_text": text,
                "style": para.style.name if para.style else None,
                "has_runs": len(para.runs) > 0,
                "runs_count": len(para.runs)
            }

            # 检测可能的变量区（包含特定关键词）
            variable_keywords = ['姓名', '性别', '年龄', '样本', '编号', 'TMB', 'MSI',
                               '日期', '医院', '科室', '医生', '电话']
            for keyword in variable_keywords:
                if keyword in text:
                    analysis["potential_variables"].append({
                        "type": "paragraph",
                        "index": i,
                        "keyword": keyword,
                        "text": text[:100]
                    })
                    break

            analysis["paragraphs"].append(para_info)

    print(f"✅ 找到 {len(analysis['paragraphs'])} 个非空段落")

    # 分析表格
    print(f"\n📊 分析表格结构...")
    for i, table in enumerate(doc.tables):
        table_info = {
            "index": i + 1,
            "rows": len(table.rows),
            "columns": len(table.columns),
            "headers": [],
            "sample_data": [],
            "contains_variables": False,
            "variable_patterns": []
        }

        # 提取表头（第一行）
        if len(table.rows) > 0:
            header_row = table.rows[0]
            table_info["headers"] = [cell.text.strip() for cell in header_row.cells]

        # 提取前3行数据作为样本
        for j, row in enumerate(table.rows[:3]):
            row_data = [cell.text.strip() for cell in row.cells]
            table_info["sample_data"].append({
                "row": j + 1,
                "cells": row_data
            })

            # 检测变量模式
            for cell_text in row_data:
                # 检测可能的变量（包含数字、特定符号等）
                if any(char.isdigit() for char in cell_text):
                    table_info["contains_variables"] = True
                # 检测基因名称模式
                if any(gene in cell_text for gene in ['ARID1A', 'JAK1', 'ERBB4', 'TP53', 'EGFR']):
                    table_info["variable_patterns"].append("基因名称")
                # 检测变异位点模式
                if 'c.' in cell_text or 'p.' in cell_text:
                    table_info["variable_patterns"].append("变异位点")
                # 检测药物名称
                if any(drug in cell_text for drug in ['顺铂', '奥沙利铂', '5-氟尿嘧啶', '卡培他滨']):
                    table_info["variable_patterns"].append("药物名称")

        # 去重变量模式
        table_info["variable_patterns"] = list(set(table_info["variable_patterns"]))

        # 统计总数据行数（排除表头）
        table_info["data_rows"] = len(table.rows) - 1 if len(table.rows) > 0 else 0

        analysis["tables"].append(table_info)

        print(f"  表格{i+1}: {len(table.rows)}行 x {len(table.columns)}列 - {table_info['variable_patterns']}")

    print(f"✅ 找到 {len(analysis['tables'])} 个表格")

    return analysis

def identify_variable_zones(analysis):
    """识别变量区域"""

    variable_zones = {
        "single_value_variables": [],  # 单值变量
        "table_variables": [],  # 表格变量
        "static_content": []  # 静态内容
    }

    # 从潜在变量中识别单值变量
    for var in analysis["potential_variables"]:
        variable_zones["single_value_variables"].append({
            "location": f"段落{var['index']}",
            "keyword": var['keyword'],
            "context": var['text']
        })

    # 识别表格变量
    for table in analysis["tables"]:
        if table["contains_variables"]:
            variable_zones["table_variables"].append({
                "table_index": table["index"],
                "rows": table["rows"],
                "columns": table["columns"],
                "headers": table["headers"],
                "data_rows": table["data_rows"],
                "variable_patterns": table["variable_patterns"]
            })
        else:
            # 可能是静态说明表格
            variable_zones["static_content"].append({
                "type": "static_table",
                "table_index": table["index"],
                "headers": table["headers"]
            })

    return variable_zones

def main():
    if len(sys.argv) < 2:
        print("用法: python analyze_generated_report.py <docx文件路径>")
        sys.exit(1)

    docx_path = sys.argv[1]

    print("=" * 80)
    print("📊 深度分析生成的报告结构")
    print("=" * 80)
    print(f"文件: {docx_path}")

    # 分析文档结构
    analysis = analyze_docx_structure(docx_path)

    # 识别变量区
    print(f"\n🔍 识别变量区域...")
    variable_zones = identify_variable_zones(analysis)

    # 输出统计
    print(f"\n" + "=" * 80)
    print("📈 分析结果统计")
    print("=" * 80)
    print(f"文档结构:")
    print(f"  - 段落数: {analysis['metadata']['paragraphs_count']}")
    print(f"  - 表格数: {analysis['metadata']['tables_count']}")
    print(f"  - 节数: {analysis['metadata']['sections_count']}")
    print(f"\n变量区统计:")
    print(f"  - 单值变量: {len(variable_zones['single_value_variables'])} 个")
    print(f"  - 表格变量: {len(variable_zones['table_variables'])} 个")
    print(f"  - 静态内容: {len(variable_zones['static_content'])} 个")

    # 输出关键表格信息
    print(f"\n📋 关键表格详情:")
    for table_var in variable_zones["table_variables"]:
        print(f"\n  表格{table_var['table_index']}:")
        print(f"    - 尺寸: {table_var['rows']}行 x {table_var['columns']}列")
        print(f"    - 数据行: {table_var['data_rows']}")
        print(f"    - 表头: {', '.join(table_var['headers'][:5])}...")
        print(f"    - 变量模式: {', '.join(table_var['variable_patterns']) if table_var['variable_patterns'] else '通用数据'}")

    # 保存详细分析结果
    output_file = "generated_report_analysis.json"
    with open(output_file, 'w', encoding='utf-8') as f:
        json.dump({
            "analysis": analysis,
            "variable_zones": variable_zones
        }, f, ensure_ascii=False, indent=2)

    print(f"\n✅ 详细分析结果已保存到: {output_file}")
    print("=" * 80)

    return analysis, variable_zones

if __name__ == "__main__":
    main()
