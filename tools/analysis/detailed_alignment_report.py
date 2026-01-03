#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
详细对齐报告工具

生成详细的对齐分析报告，包括：
- 段落级别对比
- 表格级别对比
- 可变/固定区域识别
- 差异详情输出
"""

from docx import Document
import json


def analyze_detailed_alignment(generated_file: str, reference_file: str):
    """生成详细对齐报告"""

    print("=" * 80)
    print("详细对齐分析报告")
    print("=" * 80)

    gen_doc = Document(generated_file)
    ref_doc = Document(reference_file)

    print(f"\n【文档基本信息】")
    print(f"  自动生成: {generated_file}")
    print(f"    - 段落数: {len(gen_doc.paragraphs)}")
    print(f"    - 表格数: {len(gen_doc.tables)}")
    print(f"\n  参考文件: {reference_file}")
    print(f"    - 段落数: {len(ref_doc.paragraphs)}")
    print(f"    - 表格数: {len(ref_doc.tables)}")

    # 1. 分析表格结构
    print(f"\n{'='*80}")
    print(f"【表格结构对比】")
    print(f"{'='*80}")

    print(f"\n自动生成报告的表格:")
    for i, table in enumerate(gen_doc.tables):
        rows = len(table.rows)
        cols = len(table.columns)

        # 获取第一行内容（通常是表头）
        first_row = []
        if rows > 0:
            first_row = [cell.text.strip()[:20] for cell in table.rows[0].cells[:5]]

        print(f"  表格 {i}: {rows}行 × {cols}列")
        print(f"         表头: {first_row}")

    print(f"\n参考报告的表格:")
    for i, table in enumerate(ref_doc.tables):
        rows = len(table.rows)
        cols = len(table.columns)

        # 获取第一行内容
        first_row = []
        if rows > 0:
            first_row = [cell.text.strip()[:20] for cell in table.rows[0].cells[:5]]

        print(f"  表格 {i}: {rows}行 × {cols}列")
        print(f"         表头: {first_row}")

    # 2. 关键段落对比
    print(f"\n{'='*80}")
    print(f"【关键段落对比】")
    print(f"{'='*80}")

    # 提取包含关键词的段落
    keywords = ["检测方法", "技术原理", "免责声明", "参考文献", "临床意义", "检测项目"]

    gen_key_paras = {}
    for para in gen_doc.paragraphs:
        text = para.text.strip()
        for kw in keywords:
            if kw in text:
                gen_key_paras[kw] = text[:200]
                break

    ref_key_paras = {}
    for para in ref_doc.paragraphs:
        text = para.text.strip()
        for kw in keywords:
            if kw in text:
                ref_key_paras[kw] = text[:200]
                break

    print(f"\n找到的关键段落:")
    all_keys = set(gen_key_paras.keys()) | set(ref_key_paras.keys())

    for kw in sorted(all_keys):
        gen_text = gen_key_paras.get(kw, "❌ 未找到")
        ref_text = ref_key_paras.get(kw, "❌ 未找到")

        match = "✅" if gen_text == ref_text else "⚠️"

        print(f"\n  {match} 【{kw}】")
        if gen_text != ref_text:
            print(f"      生成: {gen_text[:100]}...")
            print(f"      参考: {ref_text[:100]}...")

    # 3. 表格内容抽样对比
    print(f"\n{'='*80}")
    print(f"【表格内容抽样对比】(前10个表格)")
    print(f"{'='*80}")

    max_tables = min(10, len(gen_doc.tables), len(ref_doc.tables))

    for i in range(max_tables):
        gen_table = gen_doc.tables[i]
        ref_table = ref_doc.tables[i]

        gen_rows = len(gen_table.rows)
        gen_cols = len(gen_table.columns)
        ref_rows = len(ref_table.rows)
        ref_cols = len(ref_table.columns)

        print(f"\n  表格 {i}:")
        print(f"    结构: 生成={gen_rows}×{gen_cols}, 参考={ref_rows}×{ref_cols}")

        # 如果结构相同，抽样对比内容
        if gen_rows == ref_rows and gen_cols == ref_cols:
            diff_count = 0
            sample_diffs = []

            for row_idx in range(min(3, gen_rows)):  # 只检查前3行
                for col_idx in range(min(3, gen_cols)):  # 只检查前3列
                    gen_cell = gen_table.rows[row_idx].cells[col_idx].text.strip()
                    ref_cell = ref_table.rows[row_idx].cells[col_idx].text.strip()

                    if gen_cell != ref_cell:
                        diff_count += 1
                        if len(sample_diffs) < 3:  # 只记录前3个差异
                            sample_diffs.append({
                                "pos": f"({row_idx},{col_idx})",
                                "gen": gen_cell[:30],
                                "ref": ref_cell[:30],
                            })

            if diff_count == 0:
                print(f"    ✅ 内容一致 (抽样前3×3单元格)")
            else:
                print(f"    ⚠️  内容差异: {diff_count} 个单元格 (抽样)")
                for diff in sample_diffs:
                    print(f"       {diff['pos']}: 生成='{diff['gen']}' vs 参考='{diff['ref']}'")
        else:
            print(f"    ⚠️  结构不同，无法对比内容")

    # 4. 患者信息提取对比
    print(f"\n{'='*80}")
    print(f"【患者信息提取对比】")
    print(f"{'='*80}")

    def extract_patient_info(doc):
        """从文档中提取患者信息"""
        info = {}
        for para in doc.paragraphs[:50]:  # 只检查前50个段落
            text = para.text.strip()

            # 匹配常见字段
            import re
            patterns = {
                "姓名": r'姓名[：:]\s*(\S+)',
                "性别": r'性别[：:]\s*(\S+)',
                "年龄": r'年龄[：:]\s*(\d+)',
                "样本编号": r'样本编号[：:]\s*(\S+)',
                "检测编号": r'检测编号[：:]\s*(\S+)',
            }

            for field, pattern in patterns.items():
                match = re.search(pattern, text)
                if match:
                    info[field] = match.group(1)

        return info

    gen_patient = extract_patient_info(gen_doc)
    ref_patient = extract_patient_info(ref_doc)

    print(f"\n  自动生成报告:")
    for field, value in gen_patient.items():
        print(f"    {field}: {value}")

    print(f"\n  参考报告:")
    for field, value in ref_patient.items():
        print(f"    {field}: {value}")

    # 5. 总结
    print(f"\n{'='*80}")
    print(f"【总结】")
    print(f"{'='*80}")

    issues = []

    # 检查表格数量差异
    table_count_diff = abs(len(gen_doc.tables) - len(ref_doc.tables))
    if table_count_diff > 0:
        issues.append(f"表格数量差异: {table_count_diff} 个 (生成:{len(gen_doc.tables)}, 参考:{len(ref_doc.tables)})")

    # 检查段落数量差异（允许小幅差异）
    para_count_diff = abs(len(gen_doc.paragraphs) - len(ref_doc.paragraphs))
    if para_count_diff > 50:
        issues.append(f"段落数量差异较大: {para_count_diff} 个")

    # 检查关键段落缺失
    missing_keys = set(ref_key_paras.keys()) - set(gen_key_paras.keys())
    if missing_keys:
        issues.append(f"缺失关键段落: {', '.join(missing_keys)}")

    if issues:
        print(f"\n  ⚠️  发现的问题:")
        for issue in issues:
            print(f"    - {issue}")
    else:
        print(f"\n  ✅ 未发现明显问题")

    print(f"\n  💡 说明:")
    print(f"    - 表格数量/结构差异可能是因为参考报告和当前样本的数据不同")
    print(f"    - 患者信息差异是预期的（这是可变区域）")
    print(f"    - 关键段落（如免责声明）应该保持一致")

    print(f"\n{'='*80}")
    print(f"✅ 详细对齐分析完成")
    print(f"{'='*80}")


if __name__ == "__main__":
    generated_file = "data/output/张三---结直肠癌358基因检测-MLB2509307001-终版.docx"
    reference_file = "docs/samples/tempe_test.sanitized.docx"

    analyze_detailed_alignment(generated_file, reference_file)
