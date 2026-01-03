#!/usr/bin/env python3
"""
Highlight Jinja2 variables in a docx template.

This script creates a copy of the template with all Jinja2 variables
highlighted using different colors:
- {{ variable }} - Yellow highlight
- {% for/endfor/if/endif %} - Cyan highlight

Usage:
    python tools/highlight_jinja2_vars.py templates/jinja2_template_358_v9.docx
    python tools/highlight_jinja2_vars.py --list templates/jinja2_template_358_v9.docx

Or via CLI:
    reportgen highlight -t templates/jinja2_template_358_v9.docx
"""

import sys
import re
import argparse
from pathlib import Path
from typing import List, Tuple, Set
from copy import deepcopy

from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.table import Table
from docx.text.paragraph import Paragraph


def add_highlight_to_run(run, color: str = "yellow"):
    """
    Add highlight color to a run element.

    Args:
        run: A python-docx Run object
        color: Highlight color name (yellow, cyan, green, magenta, red, blue, etc.)
    """
    # Get or create run properties
    rPr = run._r.get_or_add_rPr()

    # Remove existing highlight if present
    existing_highlight = rPr.find(qn('w:highlight'))
    if existing_highlight is not None:
        rPr.remove(existing_highlight)

    # Add new highlight
    highlight = OxmlElement('w:highlight')
    highlight.set(qn('w:val'), color)
    rPr.append(highlight)


def contains_jinja2_variable(text: str) -> bool:
    """Check if text contains Jinja2 variable syntax {{ }}"""
    return '{{' in text and '}}' in text


def contains_jinja2_control(text: str) -> bool:
    """Check if text contains Jinja2 control syntax {% %}"""
    return '{%' in text and '%}' in text


def process_paragraph(paragraph: Paragraph) -> Tuple[int, int]:
    """
    Process a paragraph and highlight Jinja2 expressions.

    Returns:
        Tuple of (variable_count, control_count)
    """
    var_count = 0
    ctrl_count = 0

    for run in paragraph.runs:
        text = run.text or ""
        if contains_jinja2_control(text):
            add_highlight_to_run(run, "cyan")
            ctrl_count += 1
        elif contains_jinja2_variable(text):
            add_highlight_to_run(run, "yellow")
            var_count += 1

    return var_count, ctrl_count


def process_table(table: Table) -> Tuple[int, int]:
    """
    Process a table and highlight Jinja2 expressions in all cells.

    Returns:
        Tuple of (variable_count, control_count)
    """
    var_count = 0
    ctrl_count = 0

    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                v, c = process_paragraph(paragraph)
                var_count += v
                ctrl_count += c
            # Handle nested tables
            for nested_table in cell.tables:
                v, c = process_table(nested_table)
                var_count += v
                ctrl_count += c

    return var_count, ctrl_count


def highlight_jinja2_variables(input_docx: str, output_docx: str = None) -> str:
    """
    Create a highlighted version of the template.

    Args:
        input_docx: Path to input template
        output_docx: Path to output (default: input_highlighted.docx)

    Returns:
        Path to the output file
    """
    input_path = Path(input_docx)
    if output_docx is None:
        output_docx = str(input_path.parent / f"{input_path.stem}_highlighted{input_path.suffix}")

    # Load document
    doc = Document(input_docx)

    total_vars = 0
    total_ctrls = 0

    # Process main document body
    for paragraph in doc.paragraphs:
        v, c = process_paragraph(paragraph)
        total_vars += v
        total_ctrls += c

    # Process tables
    for table in doc.tables:
        v, c = process_table(table)
        total_vars += v
        total_ctrls += c

    # Process headers and footers
    for section in doc.sections:
        # Header
        for header in [section.header, section.first_page_header, section.even_page_header]:
            if header is not None:
                for paragraph in header.paragraphs:
                    v, c = process_paragraph(paragraph)
                    total_vars += v
                    total_ctrls += c
                for table in header.tables:
                    v, c = process_table(table)
                    total_vars += v
                    total_ctrls += c

        # Footer
        for footer in [section.footer, section.first_page_footer, section.even_page_footer]:
            if footer is not None:
                for paragraph in footer.paragraphs:
                    v, c = process_paragraph(paragraph)
                    total_vars += v
                    total_ctrls += c
                for table in footer.tables:
                    v, c = process_table(table)
                    total_vars += v
                    total_ctrls += c

    # Save document
    doc.save(output_docx)

    print(f"Highlighted {total_vars} variable expressions (yellow)")
    print(f"Highlighted {total_ctrls} control expressions (cyan)")
    print(f"Output saved to: {output_docx}")

    return output_docx


def extract_jinja2_expressions(input_docx: str) -> Tuple[Set[str], Set[str]]:
    """
    Extract all Jinja2 expressions from a docx template.

    Args:
        input_docx: Path to input template

    Returns:
        Tuple of (variable_expressions, control_expressions)
    """
    doc = Document(input_docx)
    variables: Set[str] = set()
    controls: Set[str] = set()

    def extract_from_text(text: str):
        """Extract Jinja2 expressions from text."""
        # Extract {{ ... }}
        for match in re.finditer(r'\{\{[^}]+\}\}', text):
            variables.add(match.group(0))
        # Extract {% ... %}
        for match in re.finditer(r'\{%[^%]+%\}', text):
            controls.add(match.group(0))

    def process_paragraphs(paragraphs):
        for paragraph in paragraphs:
            for run in paragraph.runs:
                if run.text:
                    extract_from_text(run.text)

    def process_tables(tables):
        for table in tables:
            for row in table.rows:
                for cell in row.cells:
                    process_paragraphs(cell.paragraphs)
                    process_tables(cell.tables)

    # Process main body
    process_paragraphs(doc.paragraphs)
    process_tables(doc.tables)

    # Process headers and footers
    for section in doc.sections:
        for header in [section.header, section.first_page_header, section.even_page_header]:
            if header is not None:
                process_paragraphs(header.paragraphs)
                process_tables(header.tables)
        for footer in [section.footer, section.first_page_footer, section.even_page_footer]:
            if footer is not None:
                process_paragraphs(footer.paragraphs)
                process_tables(footer.tables)

    return variables, controls


def list_jinja2_variables(input_docx: str) -> Tuple[List[str], List[str]]:
    """
    List all Jinja2 variables in the template.

    Args:
        input_docx: Path to input template

    Returns:
        Tuple of (sorted variable list, sorted control list)
    """
    variables, controls = extract_jinja2_expressions(input_docx)

    variables_list = sorted(variables)
    controls_list = sorted(controls)

    print("=" * 60)
    print(f"Jinja2 Variables ({{{{ ... }}}}) - {len(variables_list)} found:")
    print("=" * 60)
    for v in variables_list:
        print(f"  {v}")

    print()
    print("=" * 60)
    print(f"Jinja2 Control Structures ({{% ... %}}) - {len(controls_list)} found:")
    print("=" * 60)
    for c in controls_list:
        print(f"  {c}")

    return variables_list, controls_list


def main():
    """Main entry point for command-line usage."""
    parser = argparse.ArgumentParser(
        description="Highlight Jinja2 variables in a docx template"
    )
    parser.add_argument(
        "input",
        help="Path to input docx template"
    )
    parser.add_argument(
        "-o", "--output",
        help="Path to output file (default: input_highlighted.docx)"
    )
    parser.add_argument(
        "--list",
        action="store_true",
        help="List all Jinja2 expressions without creating highlighted file"
    )

    args = parser.parse_args()

    if not Path(args.input).exists():
        print(f"Error: File not found: {args.input}", file=sys.stderr)
        sys.exit(1)

    if args.list:
        list_jinja2_variables(args.input)
    else:
        highlight_jinja2_variables(args.input, args.output)


if __name__ == '__main__':
    main()
